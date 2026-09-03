#!/usr/bin/env python3
"""Normalize slide images and export crash-recoverable PDF/PPTX artifacts.

Install runtime dependencies with ``python3 -m pip install -r requirements.txt``.
"""

import argparse
import json
import os
import stat
import sys
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterable, List, Optional, Sequence

from PIL import Image
from pptx import Presentation
from pptx.util import Emu
from image_output import (
    ImageOutputError,
    ParentIdentity,
    capture_path_base,
    load_image,
    parse_json_response,
    prepare_target,
    prepared_lock_target,
    resolve_input_path,
    resolve_output_path,
    verify_parent_identity,
)
from output_lock import OutputLockError, output_lock

TARGET_SIZE = (1920, 1080)
CREAM_BACKGROUND = (248, 245, 240)
SLIDE_H = Emu(6_858_000)  # exact 7.5 inches
SLIDE_W = Emu(int(SLIDE_H) * 16 // 9)
JOURNAL_VERSION = 1
JOURNAL_SUFFIX = ".ai-image-to-ppt-export-journal.json"
COMMIT_MARKER_SUFFIX = ".commit"
PREPARATION_SUFFIX = ".prepare"
PREPARATION_OWNER_SUFFIX = ".owner"
PREPARATION_OWNER_ANCHOR_SUFFIX = ".owner.anchor"
JOURNAL_WRITE_NAME = "rollback.journal-write"
COMMIT_WRITE_NAME = "commit.journal-write"
MAX_JOURNAL_BYTES = 64 * 1024
MAX_ERROR_SUMMARY = 300
MAX_DECK_SLIDES = 128
MAX_DECK_SOURCE_BYTES = 512 * 1024 * 1024


@dataclass(frozen=True)
class LoadedJournal:
    state: object
    identity: tuple


def _flatten_transparency(image: Image.Image) -> Image.Image:
    """Return RGB pixels, compositing transparency on the deck's cream canvas."""
    has_transparency = image.mode in {"RGBA", "LA"} or (
        image.mode == "P" and "transparency" in image.info
    )
    if not has_transparency:
        return image.convert("RGB") if image.mode != "RGB" else image.copy()
    rgba = image.convert("RGBA")
    background = Image.new("RGBA", rgba.size, CREAM_BACKGROUND + (255,))
    return Image.alpha_composite(background, rgba).convert("RGB")


def normalize(image: Image.Image) -> Image.Image:
    """Cover-crop an image to an exact 1920x1080 RGB slide."""
    image = _flatten_transparency(image)
    if image.size == TARGET_SIZE:
        return image
    target_width, target_height = TARGET_SIZE
    width, height = image.size
    if width <= 0 or height <= 0:
        raise ValueError("image dimensions must be positive")
    if width * target_height > height * target_width:
        new_width = max(1, height * target_width // target_height)
        left = (width - new_width) // 2
        image = image.crop((left, 0, left + new_width, height))
    elif width * target_height < height * target_width:
        new_height = max(1, width * target_height // target_width)
        top = (height - new_height) // 2
        image = image.crop((0, top, width, top + new_height))
    return image.resize(TARGET_SIZE, Image.Resampling.LANCZOS)


def _load(path: str) -> Image.Image:
    try:
        loaded = load_image(Path(path))
        image = normalize(loaded.image)
        image.info["ai_image_to_ppt_source_bytes"] = loaded.byte_count
        return image
    except ImageOutputError as error:
        raise OSError(str(error)) from error


def _load_all(
    files: Iterable[str],
    max_total_bytes: Optional[int] = None,
) -> List[Image.Image]:
    paths = list(files)
    if not paths:
        raise ValueError("at least one input image is required")
    images = []
    total_bytes = 0
    for path in paths:
        image = _load(path)
        if max_total_bytes is not None:
            total_bytes += image.info.get("ai_image_to_ppt_source_bytes", 0)
            if total_bytes > max_total_bytes:
                raise ValueError(
                    "deck aggregate source bytes exceed maximum of "
                    f"{max_total_bytes} bytes"
                )
        images.append(image)
    return images


def _save_pdf(images: Sequence[Image.Image], path: str) -> None:
    images[0].save(
        path,
        save_all=True,
        append_images=list(images[1:]),
        resolution=150.0,
        quality=85,
        optimize=True,
    )


def _save_pptx(images: Sequence[Image.Image], path: str) -> None:
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    blank = presentation.slide_layouts[6]
    for image in images:
        buffer = BytesIO()
        image.save(buffer, format="JPEG", quality=90)
        buffer.seek(0)
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            buffer,
            0,
            0,
            width=SLIDE_W,
            height=SLIDE_H,
        )
    presentation.save(path)


def _temporary_path(
    target: Path,
    parent: Optional[ParentIdentity] = None,
) -> str:
    prepared = prepare_target(target) if parent is None else None
    identity = prepared.parent if prepared is not None else parent
    path = None
    try:
        verify_parent_identity(identity)
        descriptor, path = tempfile.mkstemp(
            prefix=f".{target.name}.",
            suffix=f".tmp{target.suffix}",
            dir=str(identity.path),
        )
        os.close(descriptor)
        verify_parent_identity(identity)
        return path
    except Exception:
        if path is not None:
            try:
                os.unlink(path)
            except OSError:
                pass
        raise


def _sync_file(path: str) -> None:
    # Windows' CRT rejects fsync/_commit on a read-only descriptor (notably on
    # Python 3.14), so reopen our freshly written temporary with write access.
    with open(path, "rb+") as stream:
        os.fsync(stream.fileno())


def _warn_retained_temp(context: str, path: str, error: object) -> None:
    try:
        print(
            f"  WARN: {context}; temporary file remains at {path}: {error}",
            file=sys.stderr,
        )
    except Exception:
        pass


def _cleanup_temp(path: str, context: str) -> None:
    last_error = None
    for _attempt in range(2):
        try:
            os.unlink(path)
            return
        except FileNotFoundError:
            return
        except OSError as error:
            last_error = error
    if last_error is not None:
        _warn_retained_temp(context, path, last_error)


def _atomic_save(target_path: str, save: Callable[[str], None]) -> None:
    target = Path(target_path)
    temporary = _temporary_path(target)
    try:
        save(temporary)
        _sync_file(temporary)
        os.replace(temporary, target)
    finally:
        _cleanup_temp(temporary, "standalone export cleanup failed")


def _freeze_inputs(files: Iterable[str], base: Path) -> List[str]:
    return [str(resolve_input_path(path, base=base)) for path in files]


def _freeze_deck_inputs(files: Iterable[str], base: Path) -> List[str]:
    frozen = []
    total_bytes = 0
    try:
        iterator = iter(files)
    except TypeError as error:
        raise ValueError("input images must be an iterable") from error
    for index, path in enumerate(iterator):
        if index >= MAX_DECK_SLIDES:
            raise ValueError(
                f"deck accepts at most {MAX_DECK_SLIDES} input slides"
            )
        resolved = resolve_input_path(path, base=base)
        try:
            source_bytes = resolved.stat().st_size
        except OSError as error:
            raise OSError(f"cannot inspect input image {resolved}: {error}") from error
        total_bytes += source_bytes
        if total_bytes > MAX_DECK_SOURCE_BYTES:
            raise ValueError(
                "deck aggregate source bytes exceed maximum of "
                f"{MAX_DECK_SOURCE_BYTES} bytes"
            )
        frozen.append(str(resolved))
    if not frozen:
        raise ValueError("at least one input image is required")
    return frozen


def export_pdf(files: Sequence[str], pdf_path: str) -> None:
    """Export a normalized multi-page PDF without exposing partial output."""
    base = capture_path_base()
    target = resolve_output_path(pdf_path, base=base)
    frozen_files = _freeze_inputs(files, base)
    images = _load_all(frozen_files)
    _atomic_save(str(target), lambda path: _save_pdf(images, path))
    size_mb = target.stat().st_size / 1024 / 1024
    print(f"  PDF: {target} ({len(images)} pages, {size_mb:.1f}MB)")


def export_pptx(files: Sequence[str], pptx_path: str) -> None:
    """Export an exact 16:9 PPTX without exposing partial output."""
    base = capture_path_base()
    target = resolve_output_path(pptx_path, base=base)
    frozen_files = _freeze_inputs(files, base)
    images = _load_all(frozen_files)
    _atomic_save(str(target), lambda path: _save_pptx(images, path))
    size_mb = target.stat().st_size / 1024 / 1024
    print(f"  PPTX: {target} ({len(images)} slides, {size_mb:.1f}MB)")


def _backup_path(
    target: Path,
    parent: Optional[ParentIdentity] = None,
) -> str:
    identity = prepare_target(target).parent if parent is None else parent
    descriptor = None
    path = None
    try:
        verify_parent_identity(identity)
        descriptor, path = tempfile.mkstemp(
            prefix=f".{target.name}.",
            suffix=".backup",
            dir=str(identity.path),
        )
        os.close(descriptor)
        descriptor = None
        verify_parent_identity(identity)
        os.unlink(path)
        return path
    except Exception:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass
        if path is not None:
            try:
                os.unlink(path)
            except OSError:
                pass
        raise


def _journal_path(output_prefix: Path) -> Path:
    prefix = Path(output_prefix)
    return prefix.parent / f".{prefix.name}{JOURNAL_SUFFIX}"


def _journal_path_for_targets(targets: Sequence[Path]) -> Path:
    first = Path(targets[0])
    prefix_name = first.name[: -len(first.suffix)] if first.suffix else first.name
    return _journal_path(first.parent / prefix_name)


def _commit_marker_path(journal: Path) -> Path:
    return journal.with_name(journal.name + COMMIT_MARKER_SUFFIX)


def _preparation_path(journal: Path) -> Path:
    return journal.with_name(journal.name + PREPARATION_SUFFIX)


def _preparation_owner_paths(preparation: Path) -> tuple:
    return (
        preparation.with_name(preparation.name + PREPARATION_OWNER_SUFFIX),
        preparation.with_name(
            preparation.name + PREPARATION_OWNER_ANCHOR_SUFFIX
        ),
    )


def _file_identity(path: Path) -> Optional[tuple]:
    try:
        current = os.stat(path, follow_symlinks=False)
    except FileNotFoundError:
        return None
    return current.st_dev, current.st_ino


def _identity_list(identity: tuple) -> List[int]:
    return [int(identity[0]), int(identity[1])]


def _bounded_error(error: object) -> str:
    text = " ".join(str(error).split()) or type(error).__name__
    if len(text) > MAX_ERROR_SUMMARY:
        text = text[: MAX_ERROR_SUMMARY - 3] + "..."
    return text


def _sync_directory(path: Path) -> None:
    flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0)
    try:
        descriptor = os.open(str(path), flags)
    except OSError:
        if os.name == "nt":
            return
        raise
    try:
        os.fsync(descriptor)
    except BaseException:
        try:
            os.close(descriptor)
        except OSError:
            pass
        raise
    os.close(descriptor)


def _directory_identity(path: Path, context: str) -> tuple:
    metadata = path.lstat()
    if not stat.S_ISDIR(metadata.st_mode):
        raise OSError(f"{context} is not a real directory: {path}")
    return metadata.st_dev, metadata.st_ino


def _preparation_owner_state(
    journal: Path,
    preparation_identity: tuple,
    parent: ParentIdentity,
) -> object:
    return {
        "version": JOURNAL_VERSION,
        "journal": str(journal),
        "preparation_identity": _identity_list(preparation_identity),
        "parent_identity": [int(parent.device), int(parent.inode)],
    }


def _write_preparation_owner(
    journal: Path,
    preparation: Path,
    preparation_identity: tuple,
    parent: ParentIdentity,
) -> tuple:
    owner, anchor = _preparation_owner_paths(preparation)
    encoded = json.dumps(
        _preparation_owner_state(journal, preparation_identity, parent),
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    descriptor = None
    owner_identity = None
    try:
        verify_parent_identity(parent)
        flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
        flags |= getattr(os, "O_CLOEXEC", 0)
        flags |= getattr(os, "O_NOFOLLOW", 0)
        descriptor = os.open(str(owner), flags, 0o600)
        metadata = os.fstat(descriptor)
        owner_identity = (metadata.st_dev, metadata.st_ino)
        # Establish the identity pair before writing. A process crash during the
        # subsequent full-record fsync then leaves two names for the same inode;
        # recovery still validates the complete record before trusting it.
        os.link(owner, anchor, follow_symlinks=False)
        with os.fdopen(descriptor, "wb") as stream:
            descriptor = None
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())
        _sync_directory(parent.path)
        return owner_identity
    except BaseException:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass
        if owner_identity is not None:
            for path in (anchor, owner):
                _remove_owned_file(path, owner_identity, "preparation owner")
        raise


def _begin_preparation(journal: Path, parent: ParentIdentity) -> tuple:
    preparation = _preparation_path(journal)
    owner_paths = _preparation_owner_paths(preparation)
    verify_parent_identity(parent)
    if any(os.path.lexists(path) for path in owner_paths):
        raise OSError(
            "unfinished export preparation ownership already exists: "
            f"{next(path for path in owner_paths if os.path.lexists(path))}"
        )
    try:
        os.mkdir(preparation, 0o700)
    except FileExistsError as error:
        raise OSError(
            f"unfinished export preparation already exists: {preparation}"
        ) from error
    identity = _directory_identity(preparation, "export preparation")
    try:
        _write_preparation_owner(journal, preparation, identity, parent)
        _sync_directory(parent.path)
    except BaseException:
        try:
            if _directory_identity(preparation, "export preparation") == identity:
                os.rmdir(preparation)
        except OSError:
            pass
        raise
    return preparation, identity


def _create_prepared_temporary(
    target: Path,
    parent: ParentIdentity,
    preparation: Path,
) -> tuple:
    verify_parent_identity(parent)
    if preparation.parent != parent.path:
        raise OSError("export preparation must share the output directory")
    descriptor = None
    temporary = None
    owner = None
    try:
        descriptor, value = tempfile.mkstemp(
            prefix=f".{target.name}.",
            suffix=f".tmp{target.suffix}",
            dir=str(preparation),
        )
        temporary = Path(value)
        temporary_stat = os.fstat(descriptor)
        identity = (temporary_stat.st_dev, temporary_stat.st_ino)
        os.close(descriptor)
        descriptor = None
        owner = temporary.with_name(temporary.name + ".owner")
        os.link(temporary, owner, follow_symlinks=False)
        _sync_directory(preparation)
        return str(temporary), str(owner), identity
    except BaseException:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass
        for candidate in (owner, temporary):
            if candidate is not None:
                try:
                    os.unlink(candidate)
                except OSError:
                    pass
        raise


def _write_record(
    record: Path,
    state: object,
    parent: ParentIdentity,
    write_name: str,
) -> tuple:
    encoded = json.dumps(
        state,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    if len(encoded) > MAX_JOURNAL_BYTES:
        raise OSError("export recovery record exceeds its size limit")

    preparation = _preparation_path(
        record if not record.name.endswith(COMMIT_MARKER_SUFFIX)
        else record.with_name(record.name[: -len(COMMIT_MARKER_SUFFIX)])
    )
    writer = preparation / write_name
    owner = writer.with_name(writer.name + ".owner")
    descriptor = None
    identity = None
    installed = False
    try:
        verify_parent_identity(parent)
        flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
        flags |= getattr(os, "O_CLOEXEC", 0)
        flags |= getattr(os, "O_NOFOLLOW", 0)
        descriptor = os.open(str(writer), flags, 0o600)
        writer_stat = os.fstat(descriptor)
        identity = (writer_stat.st_dev, writer_stat.st_ino)
        os.link(writer, owner, follow_symlinks=False)
        with os.fdopen(descriptor, "wb") as stream:
            descriptor = None
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())
        verify_parent_identity(parent)
        os.link(writer, record, follow_symlinks=False)
        installed = True
        _sync_directory(parent.path)
        return identity
    except BaseException:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass
        if not installed and identity is not None:
            for path in (writer, owner):
                _remove_owned_file(path, identity, "recovery record writer")
        raise


def _write_journal(
    journal: Path,
    state: object,
    parent: ParentIdentity,
) -> tuple:
    if not isinstance(state, dict) or state.get("decision") != "rollback":
        raise ValueError("initial export journal must record rollback")
    return _write_record(journal, state, parent, JOURNAL_WRITE_NAME)


def _write_commit_marker(
    marker: Path,
    state: object,
    parent: ParentIdentity,
) -> tuple:
    """Install a complete commit record without replacing any existing name."""
    return _write_record(marker, state, parent, COMMIT_WRITE_NAME)


def _validated_identity(value: object, label: str) -> tuple:
    if (
        not isinstance(value, list)
        or len(value) != 2
        or any(isinstance(item, bool) or not isinstance(item, int) for item in value)
    ):
        raise OSError(f"invalid export recovery journal {label}")
    return value[0], value[1]


def _load_journal(
    journal: Path,
    targets: Sequence[Path],
    parent: ParentIdentity,
) -> LoadedJournal:
    descriptor = None
    try:
        flags = os.O_RDONLY | getattr(os, "O_CLOEXEC", 0)
        flags |= getattr(os, "O_NOFOLLOW", 0)
        descriptor = os.open(str(journal), flags)
        journal_stat = os.fstat(descriptor)
        if not stat.S_ISREG(journal_stat.st_mode):
            raise OSError("export recovery journal is not a regular file")
        if journal_stat.st_size > MAX_JOURNAL_BYTES:
            raise OSError("export recovery journal exceeds its size limit")
        with os.fdopen(descriptor, "rb") as stream:
            descriptor = None
            raw = stream.read(MAX_JOURNAL_BYTES + 1)
        if len(raw) > MAX_JOURNAL_BYTES:
            raise OSError("export recovery journal exceeds its size limit")
        state = parse_json_response(raw)
    except (OSError, UnicodeError, ValueError, json.JSONDecodeError) as error:
        if isinstance(error, OSError) and str(error).startswith("export recovery"):
            raise
        raise OSError(
            f"cannot read export recovery journal {journal}: {_bounded_error(error)}"
        ) from error
    finally:
        if descriptor is not None:
            os.close(descriptor)

    if not isinstance(state, dict) or state.get("version") != JOURNAL_VERSION:
        raise OSError("invalid export recovery journal version")
    if state.get("decision") not in ("rollback", "commit"):
        raise OSError("invalid export recovery journal decision")
    records = state.get("records")
    if not isinstance(records, list) or len(records) != len(targets):
        raise OSError("invalid export recovery journal records")

    base_journal = (
        journal.with_name(journal.name[: -len(COMMIT_MARKER_SUFFIX)])
        if journal.name.endswith(COMMIT_MARKER_SUFFIX)
        else journal
    )
    preparation = _preparation_path(base_journal)
    for expected_target, record in zip(targets, records):
        if not isinstance(record, dict) or record.get("target") != str(expected_target):
            raise OSError("invalid export recovery journal target")
        temporary_value = record.get("temporary")
        if not isinstance(temporary_value, str):
            raise OSError("invalid export recovery journal temporary path")
        temporary = Path(temporary_value)
        if (
            not temporary.is_absolute()
            or temporary.parent not in (parent.path, preparation)
            or not temporary.name.startswith(f".{expected_target.name}.")
            or ".tmp" not in temporary.name
        ):
            raise OSError("invalid export recovery journal temporary path")
        _validated_identity(record.get("temporary_identity"), "temporary identity")
        owner_value = record.get("owner")
        if owner_value is not None:
            owner = Path(owner_value) if isinstance(owner_value, str) else None
            if (
                owner is None
                or not owner.is_absolute()
                or owner.parent != temporary.parent
                or owner.name != temporary.name + ".owner"
            ):
                raise OSError("invalid export recovery journal temporary owner")

        existed = record.get("existed")
        if not isinstance(existed, bool):
            raise OSError("invalid export recovery journal existed flag")
        backup_value = record.get("backup")
        original_value = record.get("original_identity")
        if existed:
            backup = Path(backup_value) if isinstance(backup_value, str) else None
            if (
                backup is None
                or not backup.is_absolute()
                or backup.parent != parent.path
                or not backup.name.startswith(f".{expected_target.name}.")
                or not backup.name.endswith(".backup")
            ):
                raise OSError("invalid export recovery journal backup path")
            _validated_identity(original_value, "original identity")
        elif backup_value is not None or original_value is not None:
            raise OSError("invalid export recovery journal original state")
    return LoadedJournal(
        state=state,
        identity=(journal_stat.st_dev, journal_stat.st_ino),
    )


def _quarantine_path(path: Path) -> Path:
    descriptor = None
    quarantine = None
    try:
        descriptor, value = tempfile.mkstemp(
            prefix=".ai-image-to-ppt-recovery-",
            suffix=".quarantine",
            dir=str(path.parent),
        )
        quarantine = Path(value)
        os.close(descriptor)
        descriptor = None
        os.unlink(quarantine)
        return quarantine
    except Exception:
        if descriptor is not None:
            try:
                os.close(descriptor)
            except OSError:
                pass
        if quarantine is not None:
            try:
                os.unlink(quarantine)
            except OSError:
                pass
        raise


def _restore_quarantined_external(
    quarantine: Path,
    path: Path,
    context: str,
) -> OSError:
    try:
        os.link(quarantine, path, follow_symlinks=False)
        restoration = "restored without clobbering the path"
    except FileExistsError:
        restoration = "path was recreated; both objects preserved"
    except (OSError, TypeError, NotImplementedError) as error:
        restoration = f"no-clobber restore failed ({_bounded_error(error)})"
    return OSError(
        f"{context} ownership changed; external file {restoration}; "
        f"quarantine preserved at {quarantine}"
    )


def _take_owned_file(
    path: Path,
    expected_identity: tuple,
    context: str,
) -> tuple:
    """Atomically move one name aside, then verify what was actually moved."""
    try:
        current_identity = _file_identity(path)
    except OSError as error:
        return None, error
    if current_identity is None:
        return None, None
    if current_identity != expected_identity:
        return None, OSError(
            f"{context} ownership changed; external file preserved"
        )

    try:
        quarantine = _quarantine_path(path)
        os.rename(path, quarantine)
        _sync_directory(path.parent)
        moved_identity = _file_identity(quarantine)
    except FileNotFoundError:
        return None, None
    except OSError as error:
        return None, error

    if moved_identity != expected_identity:
        return None, _restore_quarantined_external(
            quarantine, path, context
        )
    return quarantine, None


def _remove_quarantine(path: Path) -> Optional[OSError]:
    last_error = None
    for _attempt in range(2):
        try:
            os.unlink(path)
            _sync_directory(path.parent)
            return None
        except FileNotFoundError:
            return None
        except OSError as error:
            last_error = error
    return last_error


def _remove_owned_file(
    path: Path,
    expected_identity: tuple,
    context: str,
) -> Optional[OSError]:
    quarantine, error = _take_owned_file(path, expected_identity, context)
    if error is not None or quarantine is None:
        return error
    cleanup_error = _remove_quarantine(quarantine)
    if cleanup_error is None:
        return None
    return OSError(
        f"{context} quarantine cleanup failed; owned file preserved at "
        f"{quarantine}: {_bounded_error(cleanup_error)}"
    )


def _load_preparation_owner(
    preparation: Path,
    parent: ParentIdentity,
) -> tuple:
    owner, anchor = _preparation_owner_paths(preparation)
    descriptor = None
    try:
        owner_metadata = owner.lstat()
        anchor_metadata = anchor.lstat()
        if not stat.S_ISREG(owner_metadata.st_mode) or not stat.S_ISREG(
            anchor_metadata.st_mode
        ):
            raise OSError("preparation owner pair is not regular files")
        owner_identity = (owner_metadata.st_dev, owner_metadata.st_ino)
        if owner_identity != (anchor_metadata.st_dev, anchor_metadata.st_ino):
            raise OSError("preparation owner pair identity does not match")
        flags = os.O_RDONLY | getattr(os, "O_CLOEXEC", 0)
        flags |= getattr(os, "O_NOFOLLOW", 0)
        descriptor = os.open(str(owner), flags)
        opened = os.fstat(descriptor)
        if (opened.st_dev, opened.st_ino) != owner_identity:
            raise OSError("preparation owner identity changed while opening")
        with os.fdopen(descriptor, "rb") as stream:
            descriptor = None
            raw = stream.read(MAX_JOURNAL_BYTES + 1)
        if len(raw) > MAX_JOURNAL_BYTES:
            raise OSError("preparation owner record exceeds its size limit")
        state = parse_json_response(raw)
    except FileNotFoundError as error:
        raise OSError("preparation owner pair is missing") from error
    except (ImageOutputError, OSError, ValueError) as error:
        if isinstance(error, OSError) and str(error).startswith("preparation owner"):
            raise
        raise OSError(f"cannot read preparation owner record: {_bounded_error(error)}") from error
    finally:
        if descriptor is not None:
            os.close(descriptor)

    expected_journal = preparation.with_name(
        preparation.name[: -len(PREPARATION_SUFFIX)]
    )
    if (
        not isinstance(state, dict)
        or state.get("version") != JOURNAL_VERSION
        or state.get("journal") != str(expected_journal)
        or _validated_identity(
            state.get("parent_identity"), "preparation parent identity"
        )
        != (parent.device, parent.inode)
    ):
        raise OSError("invalid preparation owner record")
    preparation_identity = _validated_identity(
        state.get("preparation_identity"), "preparation identity"
    )
    return preparation_identity, owner_identity


def _preparation_entry_base_allowed(path: Path, targets: Sequence[Path]) -> bool:
    if path.name in (JOURNAL_WRITE_NAME, COMMIT_WRITE_NAME):
        return True
    for target in targets:
        if (
            path.name.startswith(f".{target.name}.")
            and path.name.endswith(f".tmp{target.suffix}")
        ):
            return True
    return False


def _cleanup_preparation(
    preparation: Path,
    parent: ParentIdentity,
    targets: Sequence[Path],
    expected_identity: Optional[tuple] = None,
) -> List[str]:
    owner_paths = _preparation_owner_paths(preparation)
    try:
        recorded_identity, owner_identity = _load_preparation_owner(
            preparation, parent
        )
    except OSError as error:
        return [
            _recovery_error(
                "cannot verify preparation ownership; directory preserved for "
                "offline cleanup",
                preparation,
                error,
            )
        ]
    try:
        current_identity = _directory_identity(
            preparation, "export preparation"
        )
    except FileNotFoundError:
        errors = []
        for owner_path in owner_paths:
            error = _remove_owned_file(
                owner_path, owner_identity, "preparation owner"
            )
            if error is not None:
                errors.append(
                    _recovery_error(
                        "cannot remove preparation owner", owner_path, error
                    )
                )
        return errors
    except OSError as error:
        return [_recovery_error("cannot inspect preparation", preparation, error)]
    trusted_identity = (
        recorded_identity if expected_identity is None else expected_identity
    )
    if recorded_identity != trusted_identity or current_identity != trusted_identity:
        return [
            _recovery_error(
                "cannot remove preparation",
                preparation,
                OSError("preparation ownership changed; directory preserved"),
            )
        ]

    try:
        children = list(preparation.iterdir())
    except OSError as error:
        return [_recovery_error("cannot scan preparation", preparation, error)]

    metadata_by_name = {}
    for child in children:
        try:
            metadata = child.lstat()
        except FileNotFoundError:
            continue
        except OSError as error:
            return [
                _recovery_error("cannot inspect preparation file", child, error)
            ]
        if not stat.S_ISREG(metadata.st_mode):
            return [
                _recovery_error(
                    "cannot remove preparation file",
                    child,
                    OSError("unexpected non-regular preparation entry preserved"),
                )
            ]
        metadata_by_name[child.name] = metadata

    pairs = []
    for name, metadata in metadata_by_name.items():
        if name.endswith(".owner"):
            continue
        child = preparation / name
        if not _preparation_entry_base_allowed(child, targets):
            return [
                _recovery_error(
                    "cannot remove preparation file",
                    child,
                    OSError("unknown preparation entry preserved for offline cleanup"),
                )
            ]
        owner_name = name + ".owner"
        owner_metadata = metadata_by_name.get(owner_name)
        identity = (metadata.st_dev, metadata.st_ino)
        if owner_metadata is None or identity != (
            owner_metadata.st_dev,
            owner_metadata.st_ino,
        ):
            return [
                _recovery_error(
                    "cannot remove preparation file",
                    child,
                    OSError(
                        "preparation file owner identity is unavailable; "
                        "entry preserved for offline cleanup"
                    ),
                )
            ]
        pairs.append((child, preparation / owner_name, identity))

    for name in metadata_by_name:
        if name.endswith(".owner") and name[: -len(".owner")] not in metadata_by_name:
            return [
                _recovery_error(
                    "cannot remove preparation file",
                    preparation / name,
                    OSError("orphan preparation owner preserved for offline cleanup"),
                )
            ]

    try:
        if _directory_identity(preparation, "export preparation") != trusted_identity:
            raise OSError("preparation ownership changed; directory preserved")
    except OSError as error:
        return [
            _recovery_error("cannot remove preparation", preparation, error)
        ]

    errors = []
    for child, child_owner, identity in pairs:
        for path, context in (
            (child, "preparation file"),
            (child_owner, "preparation file owner"),
        ):
            error = _remove_owned_file(path, identity, context)
            if error is not None:
                errors.append(
                    _recovery_error("cannot remove preparation file", path, error)
                )
    if errors:
        return errors
    try:
        verify_parent_identity(parent)
        if _directory_identity(preparation, "export preparation") != trusted_identity:
            raise OSError("preparation ownership changed; directory preserved")
        os.rmdir(preparation)
        _sync_directory(parent.path)
    except FileNotFoundError:
        return []
    except OSError as error:
        return [_recovery_error("cannot remove preparation", preparation, error)]
    for owner_path in owner_paths:
        error = _remove_owned_file(owner_path, owner_identity, "preparation owner")
        if error is not None:
            errors.append(
                _recovery_error("cannot remove preparation owner", owner_path, error)
            )
    return errors


def _remove_journal(
    journal: Path,
    parent: ParentIdentity,
    expected_identity: tuple,
) -> None:
    verify_parent_identity(parent)
    error = _remove_owned_file(
        journal, expected_identity, "export recovery journal"
    )
    if error is not None:
        raise error


def _recovery_error(label: str, path: Path, error: object) -> str:
    return f"{label} {path}: {_bounded_error(error)}"


def _record_source(record: object) -> Path:
    owner = record.get("owner")
    return Path(owner) if isinstance(owner, str) else Path(record["temporary"])


def _link_owned_source(
    source: Path,
    target: Path,
    expected_identity: tuple,
    parent: ParentIdentity,
    context: str,
) -> None:
    source_identity = _file_identity(source)
    if source_identity != expected_identity:
        raise OSError(f"{context} ownership changed; source preserved")
    os.link(source, target, follow_symlinks=False)
    _sync_directory(parent.path)
    target_identity = _file_identity(target)
    if target_identity == expected_identity:
        return
    if target_identity is not None:
        quarantine, isolation_error = _take_owned_file(
            target, target_identity, f"untrusted {context}"
        )
        if isolation_error is not None:
            raise OSError(
                f"{context} identity changed and final-name isolation failed: "
                f"{_bounded_error(isolation_error)}"
            )
        if quarantine is not None:
            raise OSError(
                f"{context} identity changed during publication; external file "
                f"isolated at {quarantine}"
            )
    raise OSError(f"{context} identity changed during publication")


def _retired_backup_path(backup: Path) -> Path:
    return backup.with_name(backup.name + ".retired")


def _install_backup_no_clobber(
    target: Path,
    backup: Path,
    expected_identity: tuple,
    parent: ParentIdentity,
) -> None:
    try:
        if _file_identity(target) != expected_identity:
            raise OSError(
                "forced-output ownership changed before backup; external file "
                "preserved"
            )
        os.link(target, backup, follow_symlinks=False)
        _sync_directory(parent.path)
    except FileExistsError as error:
        raise OSError(
            "forced-output backup destination appeared; external file preserved "
            f"at: {backup}"
        ) from error
    backup_identity = _file_identity(backup)
    if backup_identity != expected_identity:
        raise OSError(
            "forced-output identity changed during no-clobber backup; external "
            f"object preserved at: {backup}"
        )
    removal_error = _remove_owned_file(
        target, expected_identity, "forced output target"
    )
    if removal_error is not None:
        raise OSError(
            f"cannot retire forced output target; backup preserved at: {backup}: "
            f"{_bounded_error(removal_error)}"
        )


def _retire_backup(backup: Path, expected_identity: tuple) -> Optional[OSError]:
    retired = _retired_backup_path(backup)
    try:
        backup_identity = _file_identity(backup)
        retired_identity = _file_identity(retired)
    except OSError as error:
        return error
    if retired_identity is not None and retired_identity != expected_identity:
        return OSError("retired backup ownership changed; external file preserved")
    if backup_identity is not None:
        if backup_identity != expected_identity:
            return OSError("backup ownership changed; external file preserved")
        if retired_identity is None:
            try:
                os.link(backup, retired, follow_symlinks=False)
                _sync_directory(backup.parent)
            except FileExistsError:
                try:
                    retired_identity = _file_identity(retired)
                except OSError as error:
                    return error
                if retired_identity != expected_identity:
                    return OSError(
                        "retired backup destination appeared; external file preserved"
                    )
            except (OSError, TypeError, NotImplementedError) as error:
                return error
            try:
                retired_identity = _file_identity(retired)
            except OSError as error:
                return error
            if retired_identity != expected_identity:
                return OSError(
                    "backup ownership changed while retiring; stable recovery "
                    "files preserved"
                )
        error = _remove_owned_file(backup, expected_identity, "backup")
        if error is not None:
            return error
    if retired_identity is None:
        return None
    return _remove_owned_file(retired, expected_identity, "retired backup")


def _recover_state(
    state: object,
    journal: Path,
    parent: ParentIdentity,
    journal_identity: Optional[tuple] = None,
    commit_marker: Optional[tuple] = None,
) -> List[str]:
    records = state["records"]
    decision = state["decision"]
    errors = []

    if decision == "rollback":
        for record in records:
            target = Path(record["target"])
            temporary_identity = _validated_identity(
                record["temporary_identity"], "temporary identity"
            )
            try:
                target_identity = _file_identity(target)
            except OSError as error:
                errors.append(_recovery_error("cannot inspect target", target, error))
                continue

            if record["existed"]:
                backup = Path(record["backup"])
                original_identity = _validated_identity(
                    record["original_identity"], "original identity"
                )
                try:
                    backup_identity = _file_identity(backup)
                    retired_identity = _file_identity(
                        _retired_backup_path(backup)
                    )
                except OSError as error:
                    errors.append(
                        _recovery_error("cannot inspect backup", backup, error)
                    )
                    continue

                if backup_identity is None:
                    if target_identity == original_identity:
                        if retired_identity == original_identity:
                            error = _retire_backup(backup, original_identity)
                            if error is not None:
                                errors.append(
                                    _recovery_error(
                                        "cannot remove retired backup", backup, error
                                    )
                                )
                        continue
                    errors.append(
                        _recovery_error(
                            "cannot restore target",
                            target,
                            OSError("backup is missing; external target preserved"),
                        )
                    )
                    continue
                if backup_identity != original_identity:
                    errors.append(
                        _recovery_error(
                            "cannot restore target",
                            target,
                            OSError(
                                "backup ownership changed; external file preserved; "
                                f"backup preserved at {backup}"
                            ),
                        )
                    )
                    continue
                if target_identity not in (None, temporary_identity, original_identity):
                    errors.append(
                        _recovery_error(
                            "cannot restore target",
                            target,
                            OSError(
                                "output ownership changed; external target preserved; "
                                f"backup preserved at {backup}"
                            ),
                        )
                    )
                    continue
                try:
                    if target_identity == original_identity:
                        error = _retire_backup(backup, original_identity)
                        if error is not None:
                            raise error
                    else:
                        if target_identity == temporary_identity:
                            error = _remove_owned_file(
                                target,
                                temporary_identity,
                                "published output",
                            )
                            if error is not None:
                                raise error
                        _link_owned_source(
                            backup,
                            target,
                            original_identity,
                            parent,
                            "backup",
                        )
                        error = _retire_backup(backup, original_identity)
                        if error is not None:
                            raise error
                except (OSError, TypeError, NotImplementedError) as error:
                    errors.append(
                        _recovery_error(
                            "cannot restore target",
                            target,
                            OSError(
                                f"{_bounded_error(error)}; backup preserved at {backup}"
                            ),
                        )
                    )
            elif target_identity is not None:
                if target_identity != temporary_identity:
                    errors.append(
                        _recovery_error(
                            "cannot remove target",
                            target,
                            OSError("output ownership changed; external target preserved"),
                        )
                    )
                    continue
                error = _remove_owned_file(
                    target, temporary_identity, "published output"
                )
                if error is not None:
                    errors.append(_recovery_error("cannot remove target", target, error))
    else:
        for record in records:
            target = Path(record["target"])
            temporary = Path(record["temporary"])
            source = _record_source(record)
            temporary_identity = _validated_identity(
                record["temporary_identity"], "temporary identity"
            )
            try:
                target_identity = _file_identity(target)
                temporary_current = _file_identity(source)
            except OSError as error:
                errors.append(_recovery_error("cannot inspect committed output", target, error))
                continue
            if target_identity == temporary_identity:
                continue
            if target_identity is not None:
                errors.append(
                    _recovery_error(
                        "cannot complete committed output",
                        target,
                        OSError("output ownership changed; external target preserved"),
                    )
                )
                continue
            if temporary_current != temporary_identity:
                errors.append(
                    _recovery_error(
                        "cannot complete committed output",
                        target,
                        OSError("validated temporary file is unavailable"),
                    )
                )
                continue
            try:
                _link_owned_source(
                    source,
                    target,
                    temporary_identity,
                    parent,
                    "committed output",
                )
            except OSError as error:
                errors.append(
                    _recovery_error("cannot complete committed output", target, error)
                )

    if errors:
        return errors

    cleanup_errors = []
    for record in records:
        temporary = Path(record["temporary"])
        temporary_identity = _validated_identity(
            record["temporary_identity"], "temporary identity"
        )
        error = _remove_owned_file(
            temporary, temporary_identity, "temporary output"
        )
        if error is not None:
            cleanup_errors.append(
                _recovery_error("cannot remove temporary", temporary, error)
            )
        owner_value = record.get("owner")
        if isinstance(owner_value, str):
            owner = Path(owner_value)
            error = _remove_owned_file(
                owner, temporary_identity, "temporary owner"
            )
            if error is not None:
                cleanup_errors.append(
                    _recovery_error("cannot remove temporary owner", owner, error)
                )
        if record["existed"]:
            backup = Path(record["backup"])
            original_identity = _validated_identity(
                record["original_identity"], "original identity"
            )
            error = _retire_backup(backup, original_identity)
            if error is not None:
                cleanup_errors.append(
                    _recovery_error("cannot remove backup", backup, error)
                )

    if cleanup_errors:
        return cleanup_errors
    if journal_identity is not None:
        try:
            _remove_journal(journal, parent, journal_identity)
        except OSError as error:
            return [_recovery_error("cannot remove journal", journal, error)]
    elif os.path.lexists(journal):
        return [
            _recovery_error(
                "cannot remove journal",
                journal,
                OSError("external recovery journal appeared; file preserved"),
            )
        ]

    if commit_marker is not None:
        marker, marker_identity = commit_marker
        try:
            verify_parent_identity(parent)
            marker_error = _remove_owned_file(
                marker,
                marker_identity,
                "export commit marker",
            )
            if marker_error is not None:
                raise marker_error
        except OSError as error:
            return [
                _recovery_error("cannot remove commit marker", marker, error)
            ]
    preparation_errors = _cleanup_preparation(
        _preparation_path(journal),
        parent,
        [Path(record["target"]) for record in records],
    )
    if preparation_errors:
        return preparation_errors
    return []


def _recover_transaction(
    output_prefix: Path,
    targets: Sequence[Path],
    parent: ParentIdentity,
) -> None:
    journal = _journal_path(output_prefix)
    marker = _commit_marker_path(journal)
    preparation = _preparation_path(journal)
    journal_exists = os.path.lexists(journal)
    marker_exists = os.path.lexists(marker)
    preparation_exists = os.path.lexists(preparation)
    owner_paths = _preparation_owner_paths(preparation)
    owner_exists = any(os.path.lexists(path) for path in owner_paths)
    if not journal_exists and not marker_exists:
        if preparation_exists or owner_exists:
            errors = _cleanup_preparation(preparation, parent, targets)
            if errors:
                raise OSError(
                    "unfinished export preparation cleanup was incomplete; "
                    + "; ".join(errors)
                )
        return

    marker_record = None
    if marker_exists:
        loaded_marker = _load_journal(marker, targets, parent)
        if loaded_marker.state.get("decision") != "commit":
            raise OSError("invalid export commit marker decision")
        marker_record = (marker, loaded_marker.identity)
        state = loaded_marker.state
        expected_journal_identity = _validated_identity(
            state.get("rollback_journal_identity"),
            "rollback journal identity",
        )
    else:
        state = None
        expected_journal_identity = None

    journal_identity = None
    if journal_exists:
        if marker_record is not None:
            # A durable commit marker contains the complete recovery state and
            # binds it to the rollback journal inode. Once it exists, the commit
            # decision is authoritative even if the older rollback JSON was
            # truncated after the marker was installed.
            journal_identity = _file_identity(journal)
            if journal_identity != expected_journal_identity:
                raise OSError(
                    "export recovery journal ownership changed; external file "
                    "and commit marker preserved"
                )
        else:
            try:
                loaded = _load_journal(journal, targets, parent)
            except OSError:
                writer = preparation / JOURNAL_WRITE_NAME
                if (
                    not preparation_exists
                    or _file_identity(writer) != _file_identity(journal)
                ):
                    raise
                journal_identity = _file_identity(journal)
                error = _remove_owned_file(
                    journal, journal_identity, "incomplete export recovery journal"
                )
                if error is not None:
                    raise error
                errors = _cleanup_preparation(preparation, parent, targets)
                if errors:
                    raise OSError("; ".join(errors))
                return
            journal_identity = loaded.identity
            state = loaded.state

    errors = _recover_state(
        state,
        journal,
        parent,
        journal_identity=journal_identity,
        commit_marker=marker_record,
    )
    if errors:
        raise OSError(
            "unfinished export recovery was incomplete; " + "; ".join(errors)
        )


def _build_transaction_state(
    temporary_paths: Sequence[str],
    targets: Sequence[Path],
    force: bool,
    parent: ParentIdentity,
    owner_paths: Optional[Sequence[str]] = None,
) -> object:
    if not isinstance(force, bool):
        raise ValueError("force must be a boolean")
    if len(temporary_paths) != len(targets) or not targets:
        raise ValueError("temporary and target pairs must be non-empty and aligned")
    if owner_paths is not None and len(owner_paths) != len(targets):
        raise ValueError("temporary owner and target pairs must be aligned")
    verify_parent_identity(parent)
    targets = tuple(Path(target) for target in targets)
    if any(target.parent != parent.path for target in targets):
        raise OSError("all paired outputs must share the captured parent directory")

    existing = [target for target in targets if os.path.lexists(target)]
    if existing and not force:
        names = ", ".join(str(target) for target in existing)
        raise FileExistsError(f"output already exists; use --force: {names}")

    records = []
    owners = owner_paths if owner_paths is not None else [None] * len(targets)
    preparation = _preparation_path(_journal_path_for_targets(targets))
    for temporary_value, owner_value, target in zip(temporary_paths, owners, targets):
        temporary = Path(temporary_value)
        if temporary.parent not in (parent.path, preparation):
            raise OSError("paired temporary files must share the output directory")
        temporary_stat = os.stat(temporary, follow_symlinks=False)
        if not stat.S_ISREG(temporary_stat.st_mode):
            raise OSError(f"refusing non-regular temporary output: {temporary}")
        owner = Path(owner_value) if owner_value is not None else None
        if owner is not None:
            if owner.parent != temporary.parent:
                raise OSError("temporary owner must share the preparation directory")
            owner_stat = os.stat(owner, follow_symlinks=False)
            owner_identity = (owner_stat.st_dev, owner_stat.st_ino)
            if owner_identity != (temporary_stat.st_dev, temporary_stat.st_ino):
                raise OSError("temporary owner identity does not match output")
        existed = os.path.lexists(target)
        original_identity = None
        backup = None
        if existed:
            target_stat = os.stat(target, follow_symlinks=False)
            if not stat.S_ISREG(target_stat.st_mode):
                raise OSError(f"refusing non-regular output target: {target}")
            original_identity = (target_stat.st_dev, target_stat.st_ino)
            backup = Path(_backup_path(target, parent=parent))
        records.append({
            "target": str(target),
            "temporary": str(temporary),
            "temporary_identity": _identity_list(
                (temporary_stat.st_dev, temporary_stat.st_ino)
            ),
            "owner": str(owner) if owner is not None else None,
            "existed": existed,
            "original_identity": (
                _identity_list(original_identity)
                if original_identity is not None
                else None
            ),
            "backup": str(backup) if backup is not None else None,
        })

    return {
        "version": JOURNAL_VERSION,
        "decision": "rollback",
        "records": records,
    }


def _begin_transaction(
    temporary_paths: Sequence[str],
    targets: Sequence[Path],
    force: bool,
    parent: ParentIdentity,
    owner_paths: Optional[Sequence[str]] = None,
) -> tuple:
    state = _build_transaction_state(
        temporary_paths, targets, force, parent, owner_paths=owner_paths
    )
    journal = _journal_path_for_targets(targets)
    marker = _commit_marker_path(journal)
    if os.path.lexists(journal) or os.path.lexists(marker):
        raise OSError(
            "unfinished export recovery record already exists: "
            f"{journal if os.path.lexists(journal) else marker}"
        )
    preparation = _preparation_path(journal)
    if not os.path.lexists(preparation):
        _begin_preparation(journal, parent)
    else:
        try:
            recorded_identity, _owner_identity = _load_preparation_owner(
                preparation, parent
            )
            if (
                _directory_identity(preparation, "export preparation")
                != recorded_identity
            ):
                raise OSError("preparation ownership changed")
        except OSError as error:
            raise OSError(
                "cannot verify export preparation ownership; directory "
                f"preserved for offline cleanup: {preparation}: "
                f"{_bounded_error(error)}"
            ) from error
    journal_identity = _write_journal(journal, state, parent)
    return state, journal, journal_identity


def _warn_cleanup(message: str) -> None:
    try:
        print(message, file=sys.stderr)
    except Exception:
        pass


def _finish_visible_commit(
    marker: Path,
    journal: Path,
    targets: Sequence[Path],
    parent: ParentIdentity,
    journal_identity: tuple,
) -> bool:
    if not os.path.lexists(marker):
        return False
    try:
        loaded = _load_journal(marker, targets, parent)
        state = loaded.state
        if state.get("decision") != "commit":
            return False
        expected = _validated_identity(
            state.get("rollback_journal_identity"),
            "rollback journal identity",
        )
        if expected != journal_identity:
            return False
    except (OSError, ValueError):
        return False

    cleanup_errors = _recover_state(
        state,
        journal,
        parent,
        journal_identity=journal_identity,
        commit_marker=(marker, loaded.identity),
    )
    blocking_errors = [
        error
        for error in cleanup_errors
        if error.startswith((
            "cannot inspect committed output",
            "cannot complete committed output",
        ))
    ]
    if blocking_errors:
        raise OSError(
            "committed export verification failed; " + "; ".join(blocking_errors)
        )
    for error in cleanup_errors:
        _warn_cleanup(
            f"  WARN: committed export cleanup is incomplete; {error}"
        )
    return True


def _publish_pair(
    temporary_paths: Sequence[str],
    targets: Sequence[Path],
    force: bool,
    parent: Optional[ParentIdentity] = None,
    transaction: Optional[tuple] = None,
) -> None:
    if parent is None:
        parent = prepare_target(Path(targets[0])).parent
    if transaction is None:
        state, journal, journal_identity = _begin_transaction(
            temporary_paths, targets, force, parent
        )
    else:
        state, journal, journal_identity = transaction
    records = state["records"]
    commit_marker = _commit_marker_path(journal)
    commit_marker_identity = None

    try:
        for record in records:
            if not record["existed"]:
                continue
            target = Path(record["target"])
            backup = Path(record["backup"])
            target_identity = _validated_identity(
                record["original_identity"], "original identity"
            )
            verify_parent_identity(parent)
            _install_backup_no_clobber(
                target, backup, target_identity, parent
            )
        for record in records:
            temporary = Path(record["temporary"])
            target = Path(record["target"])
            verify_parent_identity(parent)
            temporary_identity = _validated_identity(
                record["temporary_identity"], "temporary identity"
            )
            if _file_identity(temporary) != temporary_identity:
                raise OSError(
                    f"temporary output ownership changed; external file preserved: "
                    f"{temporary}"
                )
            _link_owned_source(
                _record_source(record),
                target,
                temporary_identity,
                parent,
                "published output",
            )

        committed_state = dict(state)
        committed_state["decision"] = "commit"
        committed_state["rollback_journal_identity"] = _identity_list(
            journal_identity
        )
        commit_marker_identity = _write_commit_marker(
            commit_marker,
            committed_state,
            parent,
        )
        state = committed_state
    except Exception as publication_error:
        if _finish_visible_commit(
            commit_marker,
            journal,
            targets,
            parent,
            journal_identity,
        ):
            _warn_cleanup(
                "  WARN: commit marker was installed before a later failure; "
                f"committed outputs were preserved: {_bounded_error(publication_error)}"
            )
            return
        recovery_errors = _recover_state(
            state,
            journal,
            parent,
            journal_identity=journal_identity,
        )
        if recovery_errors:
            raise OSError(
                "publication failed: "
                f"{_bounded_error(publication_error)}; "
                "rollback was incomplete; "
                + "; ".join(recovery_errors)
            ) from publication_error
        raise

    cleanup_errors = _recover_state(
        state,
        journal,
        parent,
        journal_identity=journal_identity,
        commit_marker=(commit_marker, commit_marker_identity),
    )
    blocking_errors = [
        error
        for error in cleanup_errors
        if error.startswith((
            "cannot inspect committed output",
            "cannot complete committed output",
        ))
    ]
    if blocking_errors:
        raise OSError(
            "committed export verification failed; " + "; ".join(blocking_errors)
        )
    for error in cleanup_errors:
        _warn_cleanup(
            f"  WARN: committed export cleanup is incomplete; {error}"
        )


def _export_deck_owned(
    files: Sequence[str],
    output_prefix: str,
    parent: ParentIdentity,
    force: bool = False,
) -> None:
    """Build PDF and PPTX, then publish them with crash recovery."""
    output_prefix = os.path.abspath(output_prefix)
    targets = (Path(f"{output_prefix}.pdf"), Path(f"{output_prefix}.pptx"))
    verify_parent_identity(parent)
    images = _load_all(
        files,
        max_total_bytes=MAX_DECK_SOURCE_BYTES,
    )
    temporary_paths = []
    owner_paths = []
    transaction = None
    preparation = None
    publication_started = False
    try:
        journal = _journal_path(Path(output_prefix))
        preparation = _begin_preparation(journal, parent)
        for target in targets:
            temporary, owner, _identity = _create_prepared_temporary(
                target, parent, preparation[0]
            )
            temporary_paths.append(temporary)
            owner_paths.append(owner)
        transaction = _begin_transaction(
            temporary_paths,
            targets,
            force,
            parent,
            owner_paths=owner_paths,
        )
        _save_pdf(images, temporary_paths[0])
        _sync_file(temporary_paths[0])
        _save_pptx(images, temporary_paths[1])
        _sync_file(temporary_paths[1])
        publication_started = True
        _publish_pair(
            temporary_paths,
            targets,
            force=force,
            parent=parent,
            transaction=transaction,
        )
    except Exception as preparation_error:
        if (
            not publication_started
            and preparation is not None
        ):
            try:
                _recover_transaction(Path(output_prefix), targets, parent)
            except OSError as recovery_error:
                raise OSError(
                    "export preparation failed: "
                    f"{_bounded_error(preparation_error)}; "
                    "recovery was incomplete; "
                    f"{_bounded_error(recovery_error)}"
                ) from preparation_error
        raise
    finally:
        final_journal = _journal_path(Path(output_prefix))
        if (
            not os.path.lexists(final_journal)
            and not os.path.lexists(_commit_marker_path(final_journal))
            and not os.path.lexists(_preparation_path(final_journal))
        ):
            for temporary in temporary_paths:
                _cleanup_temp(temporary, "export preparation cleanup failed")
    for label, target in zip(("PDF", "PPTX"), targets):
        verify_parent_identity(parent)
        size_mb = target.stat().st_size / 1024 / 1024
        print(f"  {label}: {target} ({len(images)} slides, {size_mb:.1f}MB)")


def export_deck(
    files: Iterable[str], output_prefix: str, force: bool = False
) -> bool:
    """Build and recoverably publish one pair; return False for expected failures."""
    if not isinstance(force, bool):
        print("  ERR: export failed: force must be a boolean", file=sys.stderr)
        return False
    try:
        base = capture_path_base()
        resolved_prefix = resolve_output_path(output_prefix, base=base)
        targets = (
            Path(f"{resolved_prefix}.pdf"),
            Path(f"{resolved_prefix}.pptx"),
        )
        prepared_pdf = prepare_target(targets[0])
        parent = prepared_pdf.parent
        lock_target = prepared_lock_target(
            prepared_pdf,
            basename=resolved_prefix.name,
        )
        with output_lock(lock_target, namespace="deck"):
            verify_parent_identity(parent)
            _recover_transaction(resolved_prefix, targets, parent)
            frozen_files = _freeze_deck_inputs(files, base)
            _export_deck_owned(
                frozen_files,
                str(resolved_prefix),
                parent,
                force=force,
            )
        return True
    except (ImageOutputError, OutputLockError, OSError, ValueError) as error:
        print(f"  ERR: export failed: {error}", file=sys.stderr)
        return False


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Normalize images and export a PDF/PPTX pair.",
        epilog=(
            "Recovery boundary: POSIX file and directory fsync cover process "
            "crashes and power-loss metadata. Windows covers process crashes "
            "only; power-loss durability is not promised."
        ),
    )
    parser.add_argument("output_prefix", help="Output path without .pdf/.pptx")
    parser.add_argument("images", nargs="+", help="Ordered slide image paths")
    parser.add_argument(
        "--force",
        action="store_true",
        help=(
            "recoverably replace an existing PDF/PPTX pair; an interrupted "
            "publication is repaired on the next export"
        ),
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = _parser().parse_args(argv)
    return 0 if export_deck(
        args.images, args.output_prefix, force=args.force
    ) else 1


if __name__ == "__main__":
    raise SystemExit(main())
